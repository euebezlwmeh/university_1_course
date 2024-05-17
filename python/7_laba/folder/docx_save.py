from docx import Document

def docx_save(name, hours, electriciry, coast):
        document = Document()
        table = document.add_table(rows=2, cols=4)
        table.cell(0,1).text = "Количество часов"
        table.cell(0,2).text = "Потребление электроэнергии (кВт)"
        table.cell(0,3).text = "Стоимость использования прибора (руб)"
        table.cell(1,0).text = name
        table.cell(1,1).text = str(hours)
        table.cell(1,2).text = str(electriciry)
        table.cell(1,3).text = str(coast)

        document.save("7_laba.docx")