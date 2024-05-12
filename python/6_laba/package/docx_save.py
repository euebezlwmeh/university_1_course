from docx import Document

def docx_save(name, hours, A, C):
        document = Document()
        table = document.add_table(rows=2, cols=4)
        table.cell(0,1).text = "Количество часов"
        table.cell(0,2).text = "Потребление электроэнергии"
        table.cell(0,3).text = "Стоимость использования прибора"
        table.cell(1,0).text = name
        table.cell(1,1).text = hours
        table.cell(1,2).text = A
        table.cell(1,3).text = C

        document.save("6_laba.docx")